from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs" / "assessment-system-review"
OUTPUT_PATH = OUTPUT_DIR / "psychological-assessment-system-gap-analysis.docx"
HTML_NAME = "complete-care-system.html"

BLACK = "000000"
DARK_BLUE = "243B5A"
PALE_BLUE = "EDF3F8"
LIGHT_GRAY = "D9D9D9"
MUTED = "53627A"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text: str, target: str) -> None:
    rel_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5E9C")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run_pr.extend([fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([run_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_bullet(doc: Document, text: str, level=0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)


def add_numbered(doc: Document, title: str, body: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.25
    lead = paragraph.add_run(title)
    set_run_font(lead, bold=True)
    rest = paragraph.add_run(body)
    set_run_font(rest)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (cell, text) in enumerate(zip(header.cells, headers)):
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, DARK_BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        set_run_font(run, size=9.5, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            if index == 0 and len(text) <= 10:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(text)
            set_run_font(run, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title_p_pr = title._element.get_or_add_pPr()
    title_borders = title_p_pr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_p_pr.remove(title_borders)

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_run.extend([begin, instr, separate, value, end])
    paragraph._p.append(field_run)
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)


def build() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if not (OUTPUT_DIR / HTML_NAME).exists():
        raise FileNotFoundError(f"缺少 HTML 文件：{OUTPUT_DIR / HTML_NAME}")

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(8)
    remove_paragraph_borders(title)
    title_run = title.add_run("心理诊疗系统量表问题与完整诊疗差距分析")
    set_run_font(title_run, size=24, bold=True)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    meta_run = meta.add_run("面向当前 L1 至 L5 原型的设计复盘  2026年9月10日")
    set_run_font(meta_run, size=9.5, color=MUTED)

    opening = doc.add_paragraph()
    opening.paragraph_format.space_after = Pt(10)
    lead = opening.add_run("核心结论  ")
    set_run_font(lead, bold=True)
    rest = opening.add_run(
        "当前原型已经形成安全分诊、MI评估、画像确认、鉴别诊断、疗法匹配、治疗回复和输出审核的主链路，"
        "但仍不是可独立承担完整心理诊疗责任的临床系统。当前最直接的交互问题不是量表数量，而是模型提出的量表候选过早进入规则触发，"
        "使临床访谈被问卷打断。改进方向应是访谈优先、量表辅助、风险例外，并补齐临床治理、医学排查、共享决策、疗程管理、长期随访和人工转介。"
    )
    set_run_font(rest)

    link_p = doc.add_paragraph()
    link_p.paragraph_format.space_after = Pt(14)
    link_lead = link_p.add_run("交互式流程图  ")
    set_run_font(link_lead, bold=True)
    add_hyperlink(link_p, "打开完整心理诊疗系统流程图", HTML_NAME)
    path_run = link_p.add_run("  与本 Word 文件保存在同一文件夹；如 Word 弹出安全提示，请选择允许打开本地文件。")
    set_run_font(path_run, size=9.5, color=MUTED)

    doc.add_heading("一 量表问题", level=1)
    doc.add_heading("当前现象", level=2)
    p = doc.add_paragraph(
        "患者只描述一至两轮低落、焦虑或失眠，系统就可能邀请 PHQ-9、GAD-7 或 ISI。"
        "这会在治疗联盟、主诉确认、病程和功能影响尚未充分建立时，把对话重心从理解患者转向完成问卷。"
    )

    doc.add_heading("现有实现为什么会过早触发", level=2)
    add_numbered(doc, "阶段推进过快。", "第一轮提取到症状或关键叙述后，L2 即可从参与阶段进入聚焦阶段。")
    add_numbered(doc, "量表许可门槛过低。", "模型只要输出 possible_scale_domains，规则主要检查当前阶段和是否存在 focus_topics，就可将量表置为等待同意。")
    add_numbered(doc, "量表邀请会抢占普通访谈。", "一旦进入 pending_consent，下一条回复首先解释量表并请求同意。")
    add_numbered(doc, "缺少优先级与节奏管理。", "当前没有主诉确认、证据深度、访谈追问、量表队列、拒绝保护、冷却期和复测周期等统一规则。")

    doc.add_heading("建议的优先级", level=2)
    add_table(
        doc,
        ["优先级", "任务", "执行原则"],
        [
            ["P0", "安全评估和危机追问", "发现即时风险时暂停普通访谈和一般量表；直接完成意念、意图、计划、手段和时间窗口核实。"],
            ["P1", "临床访谈", "先建立关系、确认主诉，了解症状含义、病程、频率、功能影响、背景、目标与患者自己的理解。"],
            ["P2", "量表辅助", "仅在信息条件满足、用途明确且患者同意后，选择与当前主诉最相关的一项。"],
            ["P3", "结果整合", "量表结束后回到访谈，核对分数与患者体验是否一致；不能用分数直接替代诊断。"],
        ],
        [2.0, 4.0, 10.1],
    )

    doc.add_heading("量表就绪门", level=2)
    p = doc.add_paragraph("建议同时满足以下条件后，规则层才允许邀请量表：")
    for text in [
        "已经进入聚焦阶段，并由患者明确确认当前最想处理的重点；",
        "针对该领域至少完成一次开放式追问和一次反映；",
        "已经获得目标症状，以及持续时间或频率、功能影响中的必要信息；",
        "没有未完成的安全追问，也没有其他量表正在进行；",
        "本次量表具有明确用途，例如筛查、建立严重度基线、辅助鉴别或按计划复测；",
        "患者在理解用途后明确同意。",
    ]:
        add_bullet(doc, text)
    p = doc.add_paragraph(
        "不建议只用固定对话轮数作为门槛。可设置体验底线：第一条有效患者叙述后至少先完成一次反映和一次临床追问，"
        "但最终仍以信息条件是否满足为准。"
    )

    doc.add_heading("LLM与规则的职责边界", level=2)
    add_table(
        doc,
        ["组件", "应当负责", "不应当负责"],
        [
            ["LLM", "识别症状领域；提出候选量表及证据；生成自然、简短的用途说明和邀请。", "不能仅凭关键词直接启动量表；不能计分；不能覆盖安全命令或患者拒绝。"],
            ["规则", "检查量表就绪门；排序候选；一次只激活一项；处理同意、拒绝、冷却期、题号、分值和计分。", "不能把量表总分直接当成诊断；不能要求所有量表完成后才允许进入 L3。"],
        ],
        [2.5, 7.0, 6.6],
    )

    doc.add_heading("验证标准", level=2)
    add_table(
        doc,
        ["测试场景", "期望行为"],
        [
            ["第一轮只说最近没劲", "先反映和追问，不邀请量表。"],
            ["确认低落为主诉，并说明持续一个月且影响工作", "可以解释并邀请 PHQ-9，但仍需患者同意。"],
            ["同时出现低落、焦虑和失眠", "只选择与已确认主诉最相关的一项，其他进入候选队列。"],
            ["患者拒绝量表", "继续临床访谈，本阶段不重复邀请。"],
            ["出现明确自伤风险", "安全追问抢占所有一般量表和普通访谈。"],
        ],
        [6.0, 10.1],
    )

    doc.add_heading("二 当前系统相对完整诊疗系统的差距", level=1)
    p = doc.add_paragraph(
        "以下判断基于当前仓库 README 所述主链路和已检查的 L2 量表触发实现。"
        "“已实现”表示存在可运行的软件原型，不表示已经完成临床有效性验证、法规合规或取得独立诊疗资质。"
    )

    doc.add_heading("已经形成的核心能力", level=2)
    add_table(
        doc,
        ["环节", "当前能力", "成熟度"],
        [
            ["L1", "规则与 LLM 语义安全分类、风险融合、命令执行和出口守卫；模型失败时不放行。", "核心原型已实现"],
            ["L2", "MI 四阶段、自然语言提取、量表许可与计分、证据节点、画像确认和 L2 到 L3 守卫。", "部分实现"],
            ["L3", "CLINICR 候选生成、逐项排除、独立终验、疗法三步匹配和分层治疗指令。", "核心原型已实现"],
            ["L4", "按动态指令生成治疗回复，包含共情、反馈、技术执行和巩固状态。", "部分实现"],
            ["L5", "结合上下文审核 L4 回复的安全、疗法一致性和指令一致性，并支持重写与人工复核等待。", "部分实现"],
            ["编排", "评估期运行 L1 与 L2，确认后进入 L3；治疗期运行 L1、L3 动态指令、L4 与 L5，并支持回环。", "核心原型已实现"],
        ],
        [2.2, 11.2, 2.7],
    )

    doc.add_heading("仍需补齐的关键问题", level=2)
    gaps = [
        ("临床责任与治理。", "系统没有持证专业人员的责任链、人工接管后台、会诊机制、适用人群边界和临床事件处置制度。"),
        ("知情同意与隐私治理。", "目前的 API 与费用提示不能替代完整知情同意；还需覆盖数据用途、保存期限、撤回、访问控制、审计、未成年人和跨境传输。"),
        ("综合临床与医学评估。", "MI画像和结构化字段不能完全替代精神状态检查、既往史、家族史、发展史、用药与物质史、躯体疾病排查、社会支持和必要的旁系信息。"),
        ("量表策略。", "目前已具备六类量表，但触发过早；缺少就绪门、队列、拒绝保护、冷却期、基线与复测周期，以及症状分数和功能结果的联合解释。"),
        ("诊断有效性。", "CLINICR 和 Verifier 能改善推理结构，但 Verifier 评估的是证据与逻辑质量，不等于存在现实金标准，也不能单独证明诊断正确。"),
        ("共享决策。", "疗法规则已经生成推荐，但面向患者的诊断不确定性解释、备选方案、风险收益、偏好、可及性和共同签订的治疗计划仍不完整。"),
        ("完整疗程管理。", "当前可以生成单轮治疗对话，但疗程频率、阶段目标、会谈议程、练习或作业、完成情况、治疗联盟、脱落管理、治疗终止和复发预防没有形成完整闭环。"),
        ("结局监测与调整。", "尚缺少预设复评周期、患者目标结果、功能改善、症状恶化和无反应阈值，以及触发重新评估、调整疗法或升级照护的正式规则。"),
        ("危机现实闭环。", "软件可以标记、阻断和追问，但地区化热线与医疗资源、紧急联系人、真人接管、实际转介、处置回执和安全随访仍需建设。"),
        ("审核覆盖范围。", "L5 当前主要审核 L4 治疗回复；L2 MI回复、危机提示、确认卡和部分编排提示走其他路径，尚未形成统一的全系统输出质量治理。"),
        ("长期病历与跨会话连续性。", "当前会话主要保存在内存中，终端退出后不能恢复；开发者画像 HTML 不是病历，也不能作为安全、合规的会话存档。"),
        ("真实世界验证与运维。", "还需临床专家标注评测、误报与漏报分析、群体公平性、模型漂移、真实 API 稳定性、成本、日志脱敏、认证授权、限流、备份和故障恢复。"),
        ("统一产品接口。", "当前没有统一五层 HTTP 聊天 API 和患者前端；正式前端应调用复用顶层编排器的统一后端接口，而不是调用终端循环或逐层文件。"),
    ]
    for index, (title_text, body) in enumerate(gaps, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.1)
        paragraph.paragraph_format.first_line_indent = Cm(-0.1)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.18
        lead = paragraph.add_run(f"{index}.  {title_text}")
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(body)
        set_run_font(rest)

    doc.add_heading("建议的建设顺序", level=2)
    add_table(
        doc,
        ["阶段", "优先建设内容", "完成标志"],
        [
            ["近期", "量表就绪门、拒绝保护、候选队列、复测节奏；补齐自动化场景测试。", "第一轮普通主诉不再被量表抢占；多领域只启动一项；拒绝后不重复邀请。"],
            ["中期", "统一五层后端 API、持久化会话、临床评估字段、共享决策、疗程计划与过程监测。", "可以恢复跨会话疗程，并根据目标与结局数据调整方案。"],
            ["上线前", "临床治理、隐私合规、人工复核和危机转介闭环；真实世界临床与安全评估。", "关键风险可由真人接管，所有临床输出可审计，明确产品适用边界。"],
        ],
        [2.4, 7.3, 8.2],
    )

    doc.add_heading("三 完整诊疗闭环的阅读方式", level=1)
    p = doc.add_paragraph(
        "交互式 HTML 将完整流程分为入口与同意、安全分诊、临床访谈、医学补充、量表辅助、画像确认、鉴别诊断、共享决策、治疗实施、输出审核、结局监测、方案调整和长期随访。"
        "图中的绿色表示当前原型已有较明确实现，黄色表示部分覆盖，红色表示仍需建设。点击节点可查看职责和差距，支持缩放和拖动。"
    )
    p = doc.add_paragraph()
    add_hyperlink(p, "打开交互式完整心理诊疗流程", HTML_NAME)

    doc.add_heading("参考依据", level=1)
    nice_ref = doc.add_paragraph(style="List Bullet")
    nice_ref.paragraph_format.space_after = Pt(3)
    nice_ref.paragraph_format.line_spacing = 1.15
    run = nice_ref.add_run(
        "NICE 成人抑郁症治疗与管理指南 NG222。初始评估不能只依赖症状计数，还应考虑严重程度、既往史、病程、持续时间和功能损害。"
    )
    set_run_font(run, size=9.5)
    add_hyperlink(nice_ref, "查看 NICE 建议", "https://www.nice.org.uk/guidance/NG222/chapter/recommendations")

    who_ref = doc.add_paragraph(style="List Bullet")
    who_ref.paragraph_format.space_after = Pt(3)
    who_ref.paragraph_format.line_spacing = 1.15
    run = who_ref.add_run("WHO mhGAP Intervention Guide Version 2.0。以临床决策算法组织评估、管理和随访。")
    set_run_font(run, size=9.5)
    add_hyperlink(who_ref, "查看 WHO 指南", "https://www.who.int/publications/i/item/9789241549790")

    repo_ref = doc.add_paragraph(style="List Bullet")
    repo_ref.paragraph_format.space_after = Pt(3)
    repo_ref.paragraph_format.line_spacing = 1.15
    run = repo_ref.add_run(
        "本仓库 README.md，以及 layers/layer2/session/stage_controller.py、layers/layer2/scales/engine.py 和 layers/layer2/service.py 的现有实现。"
    )
    set_run_font(run, size=9.5)

    doc.core_properties.title = "心理诊疗系统量表问题与完整诊疗差距分析"
    doc.core_properties.subject = "L1至L5心理诊疗原型设计复盘"
    doc.core_properties.author = "心理诊疗系统项目"
    doc.core_properties.keywords = "心理评估, 量表, MI, L1-L5, 完整诊疗"
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build())
