from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\develop\vscode\心理诊疗系统")
REFERENCE = Path(
    r"C:\Users\ASUSA\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-system-design\assets\reference.docx"
)
OUTPUT = ROOT / "artifacts" / "reports" / "五层心理评估系统_开发进程与Schema说明.docx"

NAVY = "0B3154"
LIGHT_BLUE = "DCEAF5"
PALE_BLUE = "F3F7FB"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
BODY_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color=GRID, size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def font_run(run, *, size: float | None = None, bold: bool | None = None, color=BLACK, code=False) -> None:
    name = CODE_FONT if code else BODY_FONT
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color) if isinstance(color, str) else color


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        font_run(r1, size=10.5, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        font_run(r2, size=10.5)
    else:
        r = p.add_run(text)
        font_run(r, size=10.5)


def add_bullets(doc: Document, items: list[str], level: int = 0) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22 + 0.18 * level)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        font_run(p.add_run("• "), size=10.2, bold=True, color=NAVY)
        font_run(p.add_run(item), size=10.2)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.20)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        font_run(p.add_run(f"{index}. "), size=10.2, bold=True, color=NAVY)
        font_run(p.add_run(item), size=10.2)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    style = "Heading 1" if level == 1 else "Heading 3"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    font_run(r, size=14 if level == 1 else 11.5, bold=True, color=BLACK)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, font_size: float = 8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    set_cant_split(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=95, bottom=95)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        font_run(r, size=9, bold=True, color=RGBColor(255, 255, 255))
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        fill = PALE_BLUE if ridx % 2 else "FFFFFF"
        for cidx, text in enumerate(row_data):
            cell = row.cells[cidx]
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cell, widths[cidx])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if cidx == 2 and len(headers) == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(text))
            font_run(r, size=font_size, code=cidx in {0, 1})
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_schema_table(doc: Document, name: str, purpose: str, source: str, rows: list[tuple[str, str, str, str]]) -> None:
    add_heading(doc, name, 2)
    add_body(doc, purpose)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r1 = p.add_run("源码位置  ")
    font_run(r1, size=9, bold=True)
    r2 = p.add_run(source)
    font_run(r2, size=9, code=True)
    add_table(doc, ["字段", "类型", "必填", "字段含义"], [list(r) for r in rows], [1.45, 1.65, 0.55, 3.95], 8.4)


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Bullet 2", "List Number"):
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        style.font.name = BODY_FONT
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)
        style.font.color.rgb = BLACK
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Heading 3"].font.size = Pt(11.5)
    doc.styles["Heading 3"].font.bold = True


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    font_run(p.add_run("五层心理评估系统"), size=25, bold=False, color=RGBColor(91, 112, 133))
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(18)
    font_run(p.add_run("开发进程与 Schema 字段说明"), size=25, bold=True)
    add_body(doc, "用于项目进度汇报、代码逻辑复核和层间数据契约查阅")
    for _ in range(5):
        doc.add_paragraph()
    meta = add_table(
        doc,
        ["状态", "维护者", "更新时间"],
        [["原型已形成五层闭环", "项目开发者", "2026年9月10日"]],
        [2.35, 2.35, 2.35],
        9.3,
    )
    meta.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.rows[0].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for c in meta.rows[1].cells:
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_table(
        doc,
        ["项目项", "说明"],
        [
            ["文档范围", "当前仓库中 L1 至 L5 的已完成代码逻辑、系统流程、验证情况和主要 Schema"],
            ["主要读者", "导师、项目组成员和后续维护者"],
            ["使用方式", "先阅读进程与流程，再按层查阅 Schema 字段释义"],
            ["完整 Schema 参考", r"D:\develop\vscode\心理诊疗系统\docs\layer-schema-reference.md"],
        ],
        [1.45, 5.6],
        9.0,
    )
    doc.add_page_break()


def add_progress_report(doc: Document) -> None:
    add_heading(doc, "一 当前结论")
    add_body(doc, "当前代码已经形成安全评估、MI 信息采集、患者确认、鉴别诊断、疗法匹配、治疗对话和输出审核组成的五层闭环原型。系统能够完成一次评估期到治疗期的转换，并在治疗期为每条患者消息重新执行安全检查和输出审核。")
    add_body(doc, "这一结果说明主要业务链路已经接通，但它仍是软件原型，不能等同于完整临床诊疗产品。量表优先级、危机资源转介、真人审核、长期疗程管理、持久化病历、统一 HTTP 接口和真实世界临床验证仍需继续建设。")

    add_heading(doc, "二 整体运行逻辑")
    add_body(doc, "系统不会让每条消息机械地依次经过五层。评估期主要运行 L1 和 L2，画像经患者确认后才进入 L3；治疗期则由 L1、L3 动态指令、L4 和 L5 组成逐轮闭环。")
    flow_rows = [
        ["1", "启动", "main.py 创建总会话并加载 API 配置"],
        ["2", "安全入口", "患者消息进入 L1；危机或守卫失败时终止自动流程"],
        ["3", "信息采集", "L2 使用 MI 访谈、辅助量表和证据节点构建患者画像"],
        ["4", "患者确认", "确认后进入 L2→L3 守卫；修改或暂停留在 L2"],
        ["5", "诊断推理", "L3 运行 CLINICR 候选生成、逐步排除和独立终验"],
        ["6", "疗法与指令", "L3 用规则匹配基础疗法、插件和患者特征微调，并生成治疗指令"],
        ["7", "治疗生成", "L4 只读取上游指令生成当前轮回复，不重新进行临床推理"],
        ["8", "独立审核", "L5 检查伤害、边界和质量；失败时要求 L4 重写"],
        ["9", "持续治疗", "通过审核的回复返回患者；下一条消息再次从 L1 开始"],
    ]
    add_table(doc, ["序号", "阶段", "实际行为"], flow_rows, [0.55, 1.25, 5.3], 9.0)

    add_heading(doc, "三 启动入口与顶层编排")
    add_body(doc, "本地运行入口是 main.py。它读取或询问 API 平台、接口地址、模型和密钥，创建 L1 至 L5 所需组件，启动终端对话，并显示患者回复、内部路由和开发者画像路径。")
    add_body(doc, "真正决定下一步调用哪一层的是 orchestration/orchestrator.py。编排器保存全局会话状态，并根据守卫结果、患者确认、诊断路由和审核结果切换阶段。")
    add_table(
        doc,
        ["路由状态", "含义"],
        [
            ["L2_INTERVIEW", "正在进行普通 MI 评估访谈"],
            ["L2_SCALE", "正在执行患者已同意的量表"],
            ["L2_DIAGNOSTIC_CONTEXT", "正在补充进入 L3 所需的鉴别信息"],
            ["AWAITING_PATIENT_CONFIRMATION", "患者画像已生成，等待确认、修改或暂停"],
            ["L3_PROCESSING", "患者画像正在进入诊断推理"],
            ["RETURN_TO_L2", "L3 发现证据不足，携带缺失项返回 L2"],
            ["L4", "已进入治疗对话阶段"],
            ["MANUAL_REVIEW", "自动审核或重写无法完成，需要人工复核"],
            ["PAUSED_BY_PATIENT", "患者主动暂停"],
            ["STOPPED", "危机、安全失败或不可恢复错误导致流程终止"],
        ],
        [2.05, 5.05],
        9.0,
    )

    add_heading(doc, "四 L1 安全评估层")
    add_numbered(doc, [
        "关键词风险检测由规则执行。",
        "越狱和提示注入检测由规则执行。",
        "语义风险分类调用真实 LLM。",
        "风险融合规则合并关键词、LLM、量表上下文和越狱结果。",
        "策略执行器生成并执行安全命令。",
        "L1 出口守卫验证安全决定是否真正执行。",
    ])
    add_body(doc, "每一条新的患者消息都必须重新经过 L1，包括已进入治疗阶段后的消息。若 L1 模型不可用、输出不合法或守卫失败，系统采用 fail closed，不允许消息进入后续层。")

    add_heading(doc, "五 L2 MI 评估层")
    add_body(doc, "L2 的核心任务是收集足够好的信息，而不是下诊断。对话按参与、聚焦、唤出和计划四个 MI 阶段平滑推进，阶段名称不会直接显示给患者。")
    add_heading(doc, "L2 中的 LLM 工作", 2)
    add_bullets(doc, [
        "识别自然语言输入类型。",
        "提取症状、情绪、认知模式、患者目标和关键叙述。",
        "把复杂自然语言转换为受约束的结构化候选字段。",
        "生成符合 MI 风格的回复。",
    ])
    add_heading(doc, "L2 中的规则工作", 2)
    add_bullets(doc, [
        "验证 L1 命令和会话消息绑定。",
        "决定 MI 阶段转换。",
        "验证患者原话并生成证据节点。",
        "控制量表触发许可、题号、分值范围、计分和阈值。",
        "判断画像信息是否充分。",
        "处理患者画像确认并执行 L2→L3 接口守卫。",
    ])
    add_body(doc, "模型提取的字段只是候选事实，必须携带 evidence_quote，并由规则确认该原话确实存在于患者消息中，之后才能写入画像。")

    add_heading(doc, "六 患者画像与确认")
    add_body(doc, "L2 患者画像保存主要症状、情绪、认知模式、关键叙述、功能影响、病程、患者目标、改变动机、量表结果和鉴别诊断所需信息。每项临床事实尽量绑定患者原话和证据节点。")
    add_numbered(doc, [
        "开发者可以查看完整患者画像 HTML。",
        "患者看到脱敏确认卡，而不是内部诊断分数。",
        "患者可以确认、要求修改或暂停。",
        "只有患者确认且 L2 出口守卫通过，画像才能进入 L3。",
    ])

    add_heading(doc, "七 L3 诊断与疗法决策层")
    add_heading(doc, "CLINICR 三步推理", 2)
    add_numbered(doc, [
        "候选诊断生成：LLM 同时考虑规定的十种障碍并生成三至五个候选，记录症状匹配数、核心证据、证据引用和候选置信度。",
        "逐步排除验证：LLM 对每个候选逐条核对 DSM-5 或 ICD-11 条目，并分为满足标准、待观察或排除。每项判断只能引用 L2 中真实存在的证据。",
        "独立终验：Verifier 只审核排名第一候选的证据充分性、逻辑连贯性、鉴别覆盖和推理跳跃，不接收候选置信度，避免复制前面的结论。",
    ])
    add_body(doc, "候选置信度和 Verifier 评分是两个独立指标。候选置信度低于 0.5 时不能进入疗法匹配；Verifier 可以路由为 pass、conditional、return_to_l2、manual_review 或 reject。")
    add_heading(doc, "疗法选择三步链", 2)
    add_numbered(doc, [
        "主诊断映射基础疗法，例如 MDD 映射 CBT 加行为激活，失眠障碍映射 CBT-I。",
        "量表阈值和症状特征触发最多两个插件。L3 只输出插件名称、依据和使用要求，不在本层调用插件 API。",
        "规则根据严重程度、改变动机、精神病性症状、活跃物质依赖、严重解离和安全风险微调策略。",
    ])
    add_heading(doc, "四段治疗指令", 2)
    add_body(doc, "治疗指令中的 L1 至 L4 是四个提示词区段，不是系统的五个层级。指令 L1 定义治疗师角色，指令 L2 保存诊断锚点，指令 L3 描述本轮目标和技术并在每轮重写，指令 L4 保存绝对禁令和安全触发条件。")

    add_heading(doc, "八 L4 治疗对话层")
    add_numbered(doc, [
        "患者消息先重新经过 L1。",
        "L3 动态指令模型根据治疗历史生成当前轮指令。",
        "规则守卫检查动态指令是否符合上游诊断和疗法边界。",
        "L4 拼接角色、诊断锚点、本轮指令、防越界要求和历史。",
        "L4 调用真实 LLM 生成患者回复。",
        "本地规则先检查回复，再将完整上下文交给 L5。",
    ])
    add_body(doc, "L4 内部阶段为共情、反馈、技术执行和巩固。L4 只能执行上游批准的诊断和治疗计划，不能自行诊断新疾病或改变主诊断。")

    add_heading(doc, "九 L5 独立输出审核层")
    add_body(doc, "L5 对 L4 的候选回复进行 harm、boundary 和 quality 三项检查。审核上下文包含患者本轮原话、历史治疗对话、L1 风险信息、L3 诊断锚点、本轮动态指令和 L4 原始回复。")
    add_table(
        doc,
        ["审核结果", "后续动作"],
        [
            ["三项全部 pass", "将回复发送给患者"],
            ["任意一项 fail", "返回失败原因和修复指令，由 L4 重写"],
            ["多次重写仍失败", "暂停自动治疗并进入 MANUAL_REVIEW"],
            ["审核输出无法解析", "按不通过处理，不默认放行"],
        ],
        [2.0, 5.1],
        9.1,
    )

    add_heading(doc, "十 规则与 LLM 的职责边界")
    add_table(
        doc,
        ["位置", "LLM 负责", "规则负责"],
        [
            ["L1", "语义风险分类", "关键词、越狱检测、风险融合、策略执行、出口守卫"],
            ["L2", "输入类型识别、语义提取、复杂字段转换、MI 回复", "阶段转换、量表、证据、充分度、患者确认、接口守卫"],
            ["L3", "候选诊断、逐条排除、Verifier、动态治疗指令", "前置契约、置信度门槛、疗法选择、指令守卫、输出验证"],
            ["L4", "依据治疗指令生成回复", "阶段状态机、本地越界检查、重写次数和人工复核路由"],
            ["L5", "结合完整上下文评估伤害、边界和质量", "输入完整性、三项全通过条件、失败时阻止放行"],
        ],
        [0.7, 3.15, 3.25],
        8.8,
    )

    add_heading(doc, "十一 已完成验证与当前边界")
    add_body(doc, "最近一次真实全局测试已经完成 L1 安全评估、L2 信息采集和 PHQ-9、患者确认、L3 CLINICR 与 Verifier、L3 动态治疗指令、L4 回复生成以及 L5 审核。最终路由保持 L4，没有进入人工复核或安全阻断。")
    add_body(doc, "该次汇报记录中的自动化测试结果为 333 passed。此数字说明受控测试覆盖的代码路径通过，不代表真实 API 长期稳定，也不证明临床有效性。")
    add_table(
        doc,
        ["尚未完成的部分", "影响"],
        [
            ["量表就绪门", "当前 L2 可能在访谈信息不足时过早邀请量表"],
            ["完整危机资源与真人转介", "高风险处置还不能形成现实世界闭环"],
            ["专业人员审核后台", "MANUAL_REVIEW 暂无真人处理界面"],
            ["病历持久化与跨会话恢复", "会话退出后不能从画像 HTML 恢复"],
            ["完整疗程管理", "作业、疗效监测、阶段复评、结案和随访仍需开发"],
            ["统一五层 HTTP 接口与患者前端", "当前完整入口主要是本地终端 main.py"],
            ["临床有效性验证", "软件测试通过不能替代临床研究和专业审核"],
            ["全输出统一审核", "L5 当前主要审核 L4 治疗回复，尚未覆盖所有系统输出"],
        ],
        [2.35, 4.75],
        9.0,
    )


def add_schema_intro(doc: Document) -> None:
    add_heading(doc, "十二 Schema 阅读说明")
    add_body(doc, "Schema 是层与层之间约定的数据结构。它规定字段名称、数据类型、必填条件和允许值，使上游输出能够被下游稳定读取。Schema 通过只说明数据格式和基本约束，不等于临床结论正确。")
    add_table(
        doc,
        ["类型写法", "中文含义", "示例"],
        [
            ["str", "字符串，也就是一段文本", "session_id 或 summary"],
            ["int", "整数", "risk_level 或 severity_level"],
            ["float", "可带小数的数值", "confidence 或 verifier_score"],
            ["bool", "布尔值，只能是 true 或 false", "crisis_flag"],
            ["list[T]", "由多个同类型项目组成的列表", "list[str]"],
            ["dict[K,V]", "键值对象，键和值分别受类型约束", "dict[str, Any]"],
            ["Literal[...]", "只能从列出的固定值中选择", "pass、warn、block"],
            ["NoneType", "字段可以为空", "str | NoneType"],
            ["Any", "允许多种 JSON 类型，仍应由业务规则限制", "structured_labels"],
            ["必填", "创建对象时必须提供", "缺少后 Schema 校验失败"],
            ["默认值", "调用方没有提供时由模型自动填入", "默认空列表或 0"],
            ["工厂生成", "每次创建对象时生成独立默认容器或时间", "空列表、空字典、时间戳"],
        ],
        [1.35, 3.65, 2.1],
        9.0,
    )
    add_body(doc, "下列内容收录每层主接口 Schema 和决定行为的关键嵌套结构。完整内部模型清单以 docs/layer-schema-reference.md 为准。只要 L1 至 L5 的 Pydantic 模型或 TypedDict 发生变化，就必须重新生成并同步该参考文档。")


def add_l1_schemas(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "十三 L1 Schema")
    add_schema_table(doc, "SafetyPassport", "L1 发送给 L2 和后续层的安全通行证。它同时保存风险判断、执行结果和守卫凭证。", "shared/schemas/safety_schema.py", [
        ("schema_version", "Literal['1.0']", "是", "接口版本。用于发现上下游是否使用了不兼容的数据格式。"),
        ("layer_id", "Literal['L1']", "是", "数据来源层，固定为 L1。"),
        ("session_id", "str", "是", "全局会话标识，用来防止不同患者会话的数据混用。"),
        ("message_id", "str", "是", "当前患者消息标识。L2 入口守卫据此确认安全结果属于本条消息。"),
        ("timestamp", "str", "是", "安全评估产生时间。"),
        ("safety_flag", "pass | warn | block", "是", "总体安全状态。pass 可继续，warn 带限制继续，block 阻止常规流程。"),
        ("risk_level", "int", "是", "风险等级，当前使用 0 至 3。数值越高，处置越严格。"),
        ("risk_types", "list[str]", "是", "识别出的风险类别，例如 self_harm。"),
        ("crisis_flag", "bool", "是", "是否进入危机处理条件。"),
        ("sensitive_topics", "list[str]", "是", "需要限制追问或谨慎处理的主题。"),
        ("summary", "str", "是", "面向开发者的安全判断摘要。"),
        ("structured_labels", "dict[str, Any]", "是", "风险分项、量表跳过列表和安全追问要求等机器可读标签。"),
        ("evidence_refs", "list[EvidenceRef]", "是", "支持风险判断的证据记录。"),
        ("confidence", "float", "是", "L1 对语义分类结果的置信度，通常在 0 至 1。"),
        ("flags", "list[str]", "是", "需要下游持续执行的控制标记。"),
        ("instruction_to_next", "str", "是", "L1 交给下一层的自然语言处置说明。"),
        ("execution", "dict[str, Any]", "是", "策略执行回执，记录路由、终止状态和实际执行的命令。"),
        ("guard", "dict[str, Any]", "是", "L1 出口守卫结果，说明是否允许调用 L2 以及违反了哪些规则。"),
    ])
    add_schema_table(doc, "EvidenceRef", "L1 风险结论使用的最小证据结构。", "shared/schemas/safety_schema.py", [
        ("source", "str", "是", "证据来源，例如规则引擎或语义模型。"),
        ("finding", "str", "是", "该来源发现了什么。"),
        ("basis", "str", "是", "为什么该发现支持当前判断。"),
    ])
    add_schema_table(doc, "SafetyModelUnavailable", "L1 语义安全模型不可用时返回的终止错误结构。", "shared/schemas/safety_schema.py", [
        ("status", "Literal['error']", "是", "固定表示错误结果。"),
        ("error_code", "L1_LLM_UNAVAILABLE", "是", "专门区分 L1 模型不可用。"),
        ("message", "str", "是", "可显示或记录的错误说明。"),
        ("terminal", "Literal[True]", "是", "明确本次流程已经终止。"),
        ("can_call_l2", "Literal[False]", "是", "明确禁止继续调用 L2。"),
        ("details", "dict[str, Any]", "是", "保存错误阶段和异常类型等调试信息。"),
        ("evidence_refs", "list[EvidenceRef]", "是", "记录规则侧已经发现的安全证据，但不用于绕过模型失败。"),
    ])


def add_l2_schemas(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "十四 L2 Schema")
    add_schema_table(doc, "ClinicalIntake", "L2 交给 L3 的患者画像主接口。它必须同时表达画像内容、证据、充分度和患者确认状态。", "layers/layer2/schemas/models.py", [
        ("layer_id", "Literal['L2']", "否", "数据来源层，默认 L2。"),
        ("timestamp", "str", "否", "画像生成或更新时间。"),
        ("session_id", "str", "是", "画像所属会话。"),
        ("current_stage", "MIStage", "是", "画像生成时所在的 MI 阶段。"),
        ("key_symptoms_extracted", "list[SymptomEvidence]", "否", "已确认的主要症状及证据。"),
        ("emotions", "list[ProfileFact]", "否", "从患者原话中提取的情绪状态。"),
        ("cognitive_patterns", "list[ProfileFact]", "否", "自动化思维、核心信念或其他认知模式。"),
        ("patient_goals", "list[str]", "否", "简化后的患者目标文本。"),
        ("patient_goal_facts", "list[ProfileFact]", "否", "带原话证据的患者目标。"),
        ("key_narratives", "list[ProfileFact]", "否", "对评估重要的关键叙述。"),
        ("functional_impairments", "list[ProfileFact]", "否", "工作、学习、人际或生活功能受到的影响。"),
        ("focus_topics", "list[str]", "否", "患者确认的当前讨论重点。"),
        ("focus_topic_facts", "list[ProfileFact]", "否", "带证据的讨论重点。"),
        ("change_talk", "list[ProfileFact]", "否", "表达改变意愿、理由或能力的语句。"),
        ("sustain_talk", "list[ProfileFact]", "否", "表达维持现状、犹豫或阻力的语句。"),
        ("motivation", "MotivationScores", "是", "重要性、信心和准备度评分及来源。"),
        ("scale_results", "dict[str, ScaleResult]", "否", "各量表的状态、逐题分数、总分和阈值。"),
        ("diagnostic_context", "DiagnosticContext", "是", "L3 鉴别诊断和疗法微调需要的背景字段。"),
        ("evidence_nodes", "list[EvidenceNode]", "否", "患者原话和量表答案形成的证据图节点。"),
        ("information_sufficiency", "insufficient | sufficient", "是", "L2 是否认为画像达到进入 L3 的最低信息要求。"),
        ("missing_dimensions", "list[str]", "否", "仍缺少的临床信息维度。"),
        ("patient_confirmation", "ConfirmationStatus", "是", "患者是否已经确认画像，或要求修改、暂停。"),
        ("safety_controls", "list[str]", "否", "从 L1 继承并持续生效的安全控制。"),
        ("flags", "list[str]", "否", "L2 产生的内部状态和异常标记。"),
        ("l2_validation", "dict[str, Any]", "否", "L2 出口守卫校验结果。"),
    ])
    add_schema_table(doc, "EvidenceNode", "统一证据节点，使 L3 能从画像字段追溯到具体患者消息或量表题目。", "layers/layer2/schemas/models.py", [
        ("evidence_id", "str", "是", "证据节点唯一标识。"),
        ("source_type", "patient_message | scale_item", "是", "证据来自患者自然语言还是量表题目。"),
        ("message_id", "str", "是", "证据所属消息。"),
        ("turn_index", "int", "是", "证据出现的对话轮次，从 1 开始。"),
        ("quote", "str", "是", "患者原话或量表回答文本。"),
        ("target_path", "str", "是", "该证据支持的画像字段路径。"),
        ("created_by", "规则枚举", "是", "说明节点由经规则核验的 LLM 候选或量表规则生成。"),
        ("parent_ids", "list[str]", "否", "上游证据节点，用于表达证据之间的关系。"),
    ])
    add_schema_table(doc, "DiagnosticContext", "进入 L3 前需要补齐的鉴别诊断背景和疗法插件条件。", "layers/layer2/schemas/models.py", [
        ("symptom_duration_days", "int | None", "否", "关键症状持续天数。"),
        ("functional_impairment", "bool | None", "否", "症状是否造成明显功能影响。"),
        ("mania_or_hypomania_history", "bool | None", "否", "是否存在躁狂或轻躁狂史。"),
        ("substance_or_medical_cause", "bool | None", "否", "症状是否可能由物质、药物或躯体疾病解释。"),
        ("differential_diagnoses_assessed", "list[str] | None", "否", "已经核对过的鉴别诊断。"),
        ("evidence_conflicts", "list[dict] | None", "否", "量表、叙述或其他证据之间尚未解决的矛盾。"),
        ("relapse_count", "int | None", "否", "既往复发次数，用于插件判断。"),
        ("rhythm_disturbance", "bool | None", "否", "是否存在明显节律紊乱。"),
        ("exposure_avoidant", "bool | None", "否", "是否对暴露治疗存在显著抵触或回避。"),
        ("ritual_behaviors", "bool | None", "否", "是否存在强迫仪式行为。"),
        ("insomnia_chief_complaint", "bool | None", "否", "失眠是否为主要诉求。"),
        ("chronic_months", "int | None", "否", "相关问题持续月数。"),
        ("interpersonal_conflicts", "int | None", "否", "重要人际冲突事件数量。"),
        ("interpersonal_impairment", "bool | None", "否", "人际功能是否受损。"),
        ("single_focus", "bool | None", "否", "问题是否单一且适合聚焦处理。"),
        ("psychotic_symptoms", "bool | None", "否", "是否存在精神病性症状信号。"),
        ("substance_active", "bool | None", "否", "是否处于活跃物质依赖期。"),
        ("severe_dissociation", "bool | None", "否", "是否存在严重解离。"),
        ("sources", "dict[str, list[str]]", "否", "每个诊断背景字段对应的证据节点列表。"),
    ])
    add_schema_table(doc, "ScaleResult", "单个量表从触发到完成或拒绝的完整状态。", "layers/layer2/schemas/models.py", [
        ("scale_name", "ScaleName", "是", "量表名称，如 PHQ-9、GAD-7 或 ISI。"),
        ("status", "ScaleStatus", "否", "未开始、待同意、进行中、已完成、拒绝或跳过。"),
        ("items", "dict[int, int]", "否", "题号与患者分值。"),
        ("score", "int | None", "否", "规则计算的量表总分。"),
        ("threshold", "str | None", "否", "总分对应的分级或阈值说明。"),
        ("evidence_ids", "list[str]", "否", "每个答案产生的证据节点。"),
        ("skip_reason", "str", "否", "量表被跳过或禁止时的原因。"),
    ])
    add_schema_table(doc, "DialogueResult", "L2 自然语言适配器每处理一条消息返回的结果。", "layers/layer2/schemas/models.py", [
        ("accepted", "bool", "是", "本条消息是否被成功理解并写入状态。"),
        ("reply", "str", "是", "显示给患者的 L2 回复。"),
        ("awaiting", "str | None", "否", "下一条消息期望的输入类型。"),
        ("l2_turn", "L2Turn | None", "否", "本轮阶段、量表触发和证据摘要。"),
        ("error_code", "str | None", "否", "失败时的机器可读错误码。"),
        ("confirmation_decision", "continue | revise | pause", "否", "患者对画像确认卡的选择。"),
    ])


def add_l3_schemas(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "十五 L3 Schema")
    add_schema_table(doc, "DiagnosticCandidate", "CLINICR 第一步生成的单个候选诊断。", "layers/layer3/diagnosis/schemas.py", [
        ("disorder", "DisorderCode", "是", "十种受支持障碍中的标准代码。"),
        ("label", "str", "是", "便于阅读的诊断名称。"),
        ("symptom_match", "SymptomMatch", "是", "患者症状与诊断标准条目的匹配数和总数。"),
        ("core_evidence_summary", "str", "是", "支持该候选的关键画像信息摘要。"),
        ("evidence_refs", "list[str]", "是", "该候选引用的 L2 证据节点。"),
        ("confidence", "float", "是", "候选生成模型给出的初步主观支持度，范围 0 至 1。"),
    ])
    add_schema_table(doc, "CriterionAssessment", "CLINICR 第二步对一条 DSM-5 或 ICD-11 标准的核对结果。", "layers/layer3/diagnosis/schemas.py", [
        ("code", "str", "是", "标准条目编号，例如 A1。"),
        ("name", "str", "是", "条目名称或简述。"),
        ("status", "met | not_met | pending", "是", "已满足、不满足或信息待确认。"),
        ("rationale", "str", "是", "为何作出该状态判断。"),
        ("evidence_refs", "list[str]", "否", "支持该条目判断的患者画像证据。"),
    ])
    add_schema_table(doc, "RewardModelAssessment", "独立 Verifier 对排名第一候选推理质量的终验结果。", "layers/layer3/diagnosis/schemas.py", [
        ("score", "float", "是", "四个质量维度综合得到的终验分数。"),
        ("dimensions", "list[VerifierDimension]", "是", "证据充分、逻辑连贯、鉴别覆盖和无推理跳跃四项评分。"),
        ("fail_reasons", "list[str]", "否", "未通过或条件通过的具体原因。"),
        ("missing_dimensions", "list[MissingDimension]", "否", "需要 L2 继续追问的字段、目的和建议问题。"),
    ])
    add_schema_table(doc, "DiagnosticReport", "L3 交给 L4 的诊断、疗法和指令主接口。只有通过 L3 输出校验的报告才能进入 L4。", "layers/layer4/schemas/l3_report.py", [
        ("layer_id", "str", "否", "来源层，默认 L3。"),
        ("timestamp", "str", "否", "诊断报告生成时间。"),
        ("diagnosis", "Diagnosis", "否", "主诊断、共病、待观察和已排除诊断。"),
        ("provisional_diagnosis", "PrimaryDiagnosis | None", "否", "条件通过时保存的暂定诊断。"),
        ("verifier_score", "float", "否", "独立终验综合分数。"),
        ("verifier_decision", "str", "否", "终验给出的文字或枚举决定。"),
        ("candidate_confidence", "float", "否", "排名第一候选在第一步得到的初步置信度。"),
        ("verifier_routing", "str", "否", "pass、conditional、return_to_l2、manual_review 或 reject。"),
        ("verifier_checks", "list[VerifierCheck]", "否", "每个终验维度的状态、分数、权重、理由和证据。"),
        ("verifier_fail_reasons", "list[str]", "否", "终验失败原因。"),
        ("recommended_therapy", "RecommendedTherapy", "否", "基础疗法、插件、微调策略和禁忌。"),
        ("severity_level", "int", "否", "严重程度的结构化等级。"),
        ("clinical_notes_for_L4", "str", "否", "L4 生成治疗回复时需要遵守的临床说明。"),
        ("summary", "str", "否", "诊断和疗法结果摘要。"),
        ("structured_labels", "StructuredLabels", "否", "疗法路由、措辞限制和风险监测标签。"),
        ("evidence_refs", "list[EvidenceRefIn]", "否", "DSM 条目、疗法和规则所引用的结构化证据。"),
        ("confidence", "float", "否", "报告层面的总体置信度。"),
        ("flags", "list[str]", "否", "需要下游注意的诊断或安全标记。"),
        ("instruction_to_next", "str", "否", "交给 L4 或回到 L2 的自然语言指令。"),
        ("missing_dimensions", "list[MissingDimension]", "否", "信息不足时需要补采的字段。"),
        ("l3_validation", "L3ValidationReceipt | None", "否", "L3 输出规则校验凭证和报告哈希。"),
        ("instructions", "dict[str, Any]", "否", "L3 生成的四段治疗指令。"),
    ])
    add_schema_table(doc, "TreatmentInstructions", "L3 与 L4 之间唯一的治疗执行接口。字段名 L1 至 L4 代表指令区段。", "shared/schemas/treatment_instructions.py", [
        ("L1", "RoleInstruction", "是", "治疗师身份、语言风格和会话边界，全会话复用。"),
        ("L2", "DiagnosticAnchor", "是", "主诊断、共病、排除项、核心信念、插件、疗法和置信度等静态锚点。"),
        ("L3", "TurnInstruction | None", "否", "本轮目标、技术、禁令、上下文和证据，每轮重新生成。"),
        ("L4", "BoundaryInstruction", "是", "绝对禁止事项和安全触发条件，每轮复用。"),
    ])
    add_schema_table(doc, "TurnInstruction", "治疗期每一轮由 L3 生成并交给 L4 的动态执行指令。", "shared/schemas/treatment_instructions.py", [
        ("goal", "str", "是", "当前这一轮的具体治疗目标。"),
        ("technique", "str", "是", "建议使用的治疗技术名称。"),
        ("forbidden", "list[str]", "否", "本轮禁止行为，例如过早给建议。"),
        ("link_previous", "str", "否", "如何衔接上一轮内容。"),
        ("plugin_guidance", "str | None", "否", "已激活插件在本轮的嵌入方式。"),
        ("force_substage", "L4Substage | None", "否", "要求 L4 使用的内部治疗子阶段。"),
        ("session_id", "str", "否", "指令所属会话。"),
        ("turn_number", "int", "否", "治疗轮次编号。"),
        ("source_message", "str", "否", "用于规划本轮指令的患者原话。"),
        ("technique_code", "str", "否", "规则可验证的标准技术代码。"),
        ("plugin_steps", "dict[str, str]", "否", "每个插件对应的可执行步骤。"),
        ("evidence", "list[SourceQuote]", "否", "本轮指令所依据的患者原话。"),
        ("observations", "list[ProgressObservation]", "否", "患者反应、作业完成或未解决问题。"),
        ("previous_goal_met", "bool", "否", "上一轮目标是否达成。"),
        ("patient_ready", "bool", "否", "患者是否准备进入更具行动性的技术。"),
        ("rationale", "str", "否", "为什么选择当前目标和技术。"),
    ])


def add_l4_schemas(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "十六 L4 Schema")
    add_schema_table(doc, "SessionRequest", "创建 L4 治疗会话的输入。", "layers/layer4/service.py", [
        ("session_id", "str | None", "否", "可复用上游会话 ID；为空时生成新 ID。"),
        ("report", "DiagnosticReport", "是", "已经通过 L3→L4 守卫的诊断与疗法报告。"),
        ("safety", "SafetyPassport", "否", "L4 需要持续遵守的 L1 风险信息。"),
    ])
    add_schema_table(doc, "TurnRequest", "L4 每轮生成治疗回复的输入。", "layers/layer4/service.py", [
        ("session_id", "str", "是", "目标治疗会话。"),
        ("user_message", "str", "是", "患者当前轮自然语言。"),
        ("turn", "TurnInstruction", "是", "L3 为本轮生成且已通过守卫的动态指令。"),
    ])
    add_schema_table(doc, "TurnOutcome", "L4 完成生成和 L5 审核后的结果。", "layers/layer4/service.py", [
        ("session_id", "str", "是", "结果所属治疗会话。"),
        ("reply", "str", "是", "最终允许发送给患者的回复，或安全兜底回复。"),
        ("l4_substage", "L4-L1 至 L4-L4", "是", "当前内部治疗阶段：共情、反馈、技术执行或巩固。"),
        ("l5_context", "L5Context", "是", "本轮送审时使用的完整上下文快照。"),
        ("summary", "StructuredSummary", "是", "供后续轮次使用的结构化摘要。"),
        ("review_history", "list[ReviewResult]", "是", "首次生成和各次重写对应的 L5 审核记录。"),
        ("fallback_used", "bool", "否", "是否因多次审核失败使用兜底并转人工复核。"),
    ])
    add_schema_table(doc, "StructuredSummary", "L4 为后续轮次和追溯生成的结构化会话摘要。", "layers/layer4/schemas/summary.py", [
        ("layer_id", "Literal['L4']", "否", "来源层，固定 L4。"),
        ("timestamp", "str", "是", "摘要生成时间。"),
        ("summary", "str", "是", "本轮治疗内容摘要。"),
        ("structured_labels", "dict[str, Any]", "是", "本轮阶段、技术和状态等机器可读标签。"),
        ("evidence_refs", "list[EvidenceRef]", "否", "摘要内容引用的证据。"),
        ("confidence", "float", "否", "摘要可靠度。"),
        ("flags", "list[str]", "否", "后续轮次需要关注的内部标记。"),
    ])


def add_l5_schemas(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "十七 L5 Schema")
    add_schema_table(doc, "L5Context", "L4 交给 L5 的完整审核上下文。L5 必须看到原始回复、患者原话、历史、指令和风险信息。", "layers/layer4/schemas/l5_context.py", [
        ("l4_raw_output", "str", "是", "L4 尚未发给患者的原始候选回复。"),
        ("l4_substage", "L4Substage", "是", "当前回复所属的 L4 内部治疗阶段。"),
        ("therapy_type", "str", "是", "L3 批准的基础疗法。"),
        ("user_message", "str", "是", "患者当前轮原话，不能为空。"),
        ("history", "list[ConversationTurn]", "是", "此前患者与系统的治疗对话。"),
        ("instructions", "TreatmentInstructions", "是", "完整四段治疗指令，包含本轮动态指令。"),
        ("diagnosis_context", "DiagnosisContext", "否", "主诊断、共病和自杀意念摘要。"),
        ("risk_context", "RiskContext", "否", "L1 风险等级、敏感主题和安全标记。"),
    ])
    add_schema_table(doc, "ReviewResult", "L5 对候选回复进行三项独立检查后的输出。", "layers/layer4/schemas/review.py", [
        ("harm", "pass | fail", "是", "回复是否可能造成伤害或弱化必要的安全处置。"),
        ("boundary", "pass | fail", "是", "回复是否越过心理支持和治疗执行边界。"),
        ("quality", "pass | fail", "是", "回复是否符合当前疗法、阶段和本轮动态指令。"),
        ("fail_reason", "str", "否", "任一检查失败时说明具体问题。"),
        ("fix_instruction", "str", "否", "交给 L4 重写模型的可执行修复要求。"),
    ])
    add_schema_table(doc, "DiagnosisContext", "L5 审核所需的最小诊断背景。", "layers/layer4/schemas/l5_context.py", [
        ("primary", "str", "否", "主诊断或初步诊断。"),
        ("secondary", "list[str]", "否", "共病或次要诊断。"),
        ("suicidal_ideation", "str | None", "否", "与当前安全审核有关的自杀意念信息。"),
    ])
    add_schema_table(doc, "RiskContext", "L5 审核所需的最小安全背景。", "layers/layer4/schemas/l5_context.py", [
        ("risk_level", "int", "否", "L1 当前风险等级。"),
        ("sensitive_topics", "list[str]", "否", "需要避免误触或谨慎回应的主题。"),
        ("flags", "list[str]", "否", "仍需执行的安全控制标记。"),
    ])


def add_orchestrator_schema(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "十八 顶层编排 Schema")
    add_schema_table(doc, "OrchestrationOutcome", "顶层编排器每处理一条消息后交给终端或未来前端的统一结果。", "orchestration/schemas.py", [
        ("session_id", "str", "是", "当前全局会话。"),
        ("route", "PipelineRoute", "是", "系统当前所在层级或等待状态。"),
        ("reply", "str", "是", "本轮可显示给患者的文字。"),
        ("waiting_for", "str | None", "否", "下一条消息期望的输入类型。"),
        ("terminal", "bool", "否", "本次会话是否已经终止。"),
        ("confirmation", "PatientConfirmation | None", "否", "信息充分时生成的患者确认卡。"),
    ])
    add_schema_table(doc, "PatientConfirmation", "患者画像完成后显示给患者的脱敏确认卡。", "orchestration/schemas.py", [
        ("title", "str", "是", "确认卡标题。"),
        ("profile_summary", "list[str]", "否", "主要症状、病程、功能影响和目标摘要。"),
        ("scale_summary", "list[str]", "否", "已完成量表及分数摘要。"),
        ("initial_impression", "str", "是", "非确诊措辞的初步整理。"),
        ("disclaimer", "str", "是", "量表和系统输出不能替代专业诊断的说明。"),
        ("next_step", "str", "是", "确认后系统将执行的动作。"),
        ("actions", "continue | revise | pause", "否", "患者可以选择继续、修改或暂停。"),
    ])

    add_heading(doc, "十九 下一步工作")
    add_table(
        doc,
        ["优先级", "工作项", "完成标准"],
        [
            ["P0", "增加量表就绪门", "在充分开放访谈和必要澄清前不主动邀请量表；高风险量表题除外"],
            ["P0", "建设危机转介闭环", "按地区展示可用资源，支持真人接管并记录处置结果"],
            ["P1", "持久化会话和病历", "安全保存状态、支持跨进程恢复并具有审计和删除机制"],
            ["P1", "统一五层 HTTP 接口", "前端只调用编排器 API，不直接逐层调用"],
            ["P1", "专业审核后台", "能够处理 MANUAL_REVIEW、记录决定并安全恢复会话"],
            ["P2", "完整疗程管理", "支持疗程目标、家庭作业、疗效复评、结案和随访"],
            ["P2", "临床与真实 API 评估", "分别验证可靠性、响应时间、证据一致性和临床安全性"],
        ],
        [0.8, 2.3, 4.0],
        8.9,
    )
    add_heading(doc, "二十 汇报摘要")
    add_body(doc, "当前项目已经完成基于顶层编排器的五层心理评估与治疗原型。评估期由 L1 完成安全筛查，L2 通过 MI 访谈和辅助量表构建可追溯的患者画像；患者确认后，L3 使用 CLINICR 三步法进行鉴别诊断，并通过规则完成疗法选择和治疗指令生成；治疗期每轮重新经过 L1 安全检查，由 L3 生成动态指令、L4 执行治疗对话、L5 进行输出审核，不通过时自动重写或转人工复核。当前下一步重点是降低量表依赖，并补齐长期疗程管理、真人转介、数据持久化和统一前端接口。")


def finalize(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "五层心理评估系统 | 开发进程与 Schema 字段说明"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        font_run(r, size=8, color=RGBColor(91, 112, 133))

    core = doc.core_properties
    core.title = "五层心理评估系统 开发进程与 Schema 字段说明"
    core.subject = "L1 至 L5 代码逻辑、系统流程与数据接口说明"
    core.author = "项目开发者"
    core.keywords = "心理评估系统,L1,L2,L3,L4,L5,Schema,CLINICR,MI"
    core.comments = ""


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)
    configure_styles(doc)
    add_cover(doc)
    add_progress_report(doc)
    add_schema_intro(doc)
    add_l1_schemas(doc)
    add_l2_schemas(doc)
    add_l3_schemas(doc)
    add_l4_schemas(doc)
    add_l5_schemas(doc)
    add_orchestrator_schema(doc)
    finalize(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
