from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\develop\vscode\心理诊疗系统")
OUTPUT = ROOT / "artifacts" / "reports" / "心理诊疗系统五项后续开发任务说明.docx"

NAVY = "244A64"
PALE_BLUE = "EDF4F8"
PALE_GRAY = "F5F6F7"
LIGHT_BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_BORDER)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_layout(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for r_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for c_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = widths[c_index]
            if r_index == 0:
                set_cell_shading(cell, NAVY)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
            elif r_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    if r_index != 0:
                        run.font.size = Pt(9.3)
                    set_run_fonts(run)


def set_run_fonts(run, east_asia="Microsoft YaHei", latin="Aptos"):
    run.font.name = latin
    run._element.get_or_add_rPr()
    rfonts = run._element.rPr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        run._element.rPr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_fonts(run)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_fonts(run2)


def remove_paragraph_borders(paragraph_or_style):
    element = paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.14
    run = p.add_run(text)
    set_run_fonts(run)
    return p


def add_number(doc, text, level=0):
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.18
    run = p.add_run(text)
    set_run_fonts(run)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead:
        run = p.add_run(bold_lead)
        run.bold = True
        set_run_fonts(run)
    run = p.add_run(text)
    set_run_fonts(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        run.font.color.rgb = BLACK
        set_run_fonts(run)
    return p


def add_task_table(doc, rows):
    table = doc.add_table(rows=1, cols=5)
    headers = ["编号", "任务", "具体工作", "交付物", "验收标准"]
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for item in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(item):
            cells[idx].text = value
    set_repeat_table_layout(
        table,
        [Inches(0.55), Inches(1.15), Inches(2.05), Inches(1.35), Inches(1.55)],
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_simple_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for item in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(item):
            cells[idx].text = value
    set_repeat_table_layout(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    remove_paragraph_borders(title)

    for name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        styles[name].font.name = "Aptos"
        styles[name].font.size = Pt(10.5)
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    add_page_number(footer_p)
    for run in footer_p.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(90, 90, 90)

    p = doc.add_paragraph(style="Title")
    remove_paragraph_borders(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run("心理诊疗系统五项后续开发任务说明")
    set_run_fonts(run, east_asia="Microsoft YaHei", latin="Aptos Display")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("前端  L2问诊知识库  L3临床知识库  Patient Agent  通用RAG检索引擎")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(55, 75, 88)
    set_run_fonts(run)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(90)
    for line in (
        "用途  团队任务分配和开发验收",
        "适用项目  五层心理诊疗大模型系统",
        f"版本日期  {date.today().isoformat()}",
    ):
        r = meta.add_run(line + "\n")
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(70, 70, 70)
        set_run_fonts(r)

    doc.add_page_break()

    add_heading(doc, "文档目的", 1)
    add_body(
        doc,
        "本文将五项后续开发工作拆分为可直接分配的工作包，明确每项工作的功能边界、实施步骤、接口、交付物、依赖关系和验收标准。现有L1至L5主流程继续保留，新增模块通过统一Schema和顶层编排器接入，不重新实现已经完成的安全分类、MI阶段控制、CLINICR、治疗指令或L5单轮审核。",
    )
    add_body(
        doc,
        "五项工作的共同目标是让系统获得可检索的问诊与临床知识、可重复的患者模拟能力，以及可以用于项目展示的统一交互界面。L2知识库负责帮助系统决定如何提问，L3知识库负责提供诊断标准依据，通用RAG只负责检索，Patient Agent只负责模拟测试患者，前端只通过完整系统API与顶层编排器交互。",
    )

    add_heading(doc, "五项工作边界", 1)
    add_simple_table(
        doc,
        ["工作包", "解决的问题", "不负责的内容", "主要依赖"],
        [
            ("通用RAG检索引擎", "统一完成知识索引 Top K检索 过滤 重排和引用返回", "不诊断 不生成患者事实 不决定MI阶段", "知识库Schema和数据"),
            ("L2问诊知识库", "提供MI话术 症状追问和量表自然引入范式", "不生成诊断标签 不替代量表规则", "通用RAG和现有L2"),
            ("L3临床知识库", "提供疾病标准 排除条件 鉴别诊断和来源版本", "不保存患者事实 不直接生成治疗回复", "通用RAG和现有CLINICR"),
            ("Patient Agent", "模拟稳定一致的患者并自动完成多轮测试", "不作为临床金标准 不进入生产患者流程", "病例Schema和完整系统API"),
            ("前端", "提供患者对话 开发者流程观察和项目演示", "不直接调用单层 不在界面内做临床判断", "顶层编排器和完整系统API"),
        ],
        [Inches(1.25), Inches(2.2), Inches(2.0), Inches(1.3)],
    )

    add_heading(doc, "总体依赖关系", 1)
    add_body(doc, "数据整理和接口设计可以并行启动，模块接入必须按以下依赖顺序推进。")
    for text in (
        "先冻结RetrievedKnowledge L2KnowledgeCard ClinicalCriterion PatientScenario等共享Schema。",
        "L2和L3知识负责人分别整理数据，RAG负责人同步搭建通用检索能力。",
        "通用RAG完成后，L2和L3负责人分别接入对应知识库。",
        "顶层编排器提供完整系统API后，Patient Agent和前端再进行端到端联调。",
        "五项工作最终进入同一套单元测试 接口测试 受控模型测试和真实LLM测试。",
    ):
        add_number(doc, text)

    add_heading(doc, "一  通用RAG检索引擎", 1)
    add_body(
        doc,
        "通用RAG是公共基础设施，负责把查询转换成可追溯的知识片段。它同时服务L2和L3，但两层使用独立知识集合、过滤规则和访问权限。RAG返回资料及来源，不进行诊断，也不能把知识片段写成患者事实。",
    )
    add_heading(doc, "功能流程", 2)
    for text in (
        "接收查询文本 collection Top K 元数据过滤条件和调用层级。",
        "根据collection选择L2问诊库或L3临床库，禁止跨库混检。",
        "执行关键词与向量混合检索，并按元数据进行精确过滤。",
        "对候选结果去重和重排，低于相关性阈值的结果不返回。",
        "返回knowledge id 内容 来源 版本 位置 分数和索引版本。",
        "记录检索日志，供开发者界面和问题追踪使用。",
    ):
        add_number(doc, text)

    add_heading(doc, "核心接口", 2)
    add_simple_table(
        doc,
        ["接口字段", "类型", "含义", "控制方式"],
        [
            ("collection", "字符串枚举", "指定L2问诊库或L3临床库", "规则白名单"),
            ("query", "字符串", "当前检索请求", "调用层构造"),
            ("top k", "整数", "最多返回的结果数量", "配置限制"),
            ("filters", "对象", "阶段 疾病代码 版本 来源类型等过滤条件", "规则生成"),
            ("results", "数组", "检索到的结构化知识条目", "Schema校验"),
            ("index version", "字符串", "本次使用的索引版本", "系统写入"),
        ],
        [Inches(1.15), Inches(1.0), Inches(2.8), Inches(1.75)],
    )

    add_heading(doc, "任务清单", 2)
    add_task_table(doc, [
        ("R1", "统一检索Schema", "定义查询 过滤条件 检索结果 来源和索引版本字段", "Pydantic模型和字段文档", "非法集合 未知字段和缺失来源会被拒绝"),
        ("R2", "知识导入", "读取标准化知识卡 校验来源 生成索引记录", "导入器和错误报告", "坏数据不能进入索引 每条记录可追溯"),
        ("R3", "索引构建", "建立关键词索引和向量索引 保存版本", "可重建索引和清单", "同一版本可重复构建且记录数一致"),
        ("R4", "混合检索", "实现元数据过滤 关键词召回 向量召回 去重和Top K", "统一retriever接口", "L2与L3集合隔离 Top K和阈值生效"),
        ("R5", "重排策略", "优先保留来源质量 高匹配度和不同用途的结果", "reranker和配置", "不返回大量重复范式或错误疾病资料"),
        ("R6", "安全控制", "把检索内容标记为数据 防止知识文本覆盖系统指令", "输入清洗和提示边界", "恶意知识条目不能改变调用层职责"),
        ("R7", "日志与观测", "记录查询 过滤条件 命中条目 分数 延迟和版本", "结构化检索日志", "能定位空结果 错库和低相关问题"),
        ("R8", "检索评测", "建立固定查询集 计算命中率 相关性和来源完整率", "评测脚本与报告", "关键查询能找到预期条目且无跨库结果"),
    ])

    add_heading(doc, "负责人交付边界", 2)
    add_bullet(doc, "建议负责人  RAG或算法工程成员。")
    add_bullet(doc, "允许修改  新建retrieval目录 shared中的检索Schema和配置。")
    add_bullet(doc, "不得修改  L2阶段规则 L3诊断门槛和L4治疗策略。")
    add_bullet(doc, "完成定义  L2和L3能通过同一接口检索各自知识，并返回可验证来源。")

    doc.add_page_break()
    add_heading(doc, "二  L2问诊知识库", 1)
    add_body(
        doc,
        "L2知识库用于改善MI访谈，不用于诊断。它向L2模型提供当前阶段适用的OARS话术、症状探索范式、量表说明和自然衔接示例。知识条目只能作为语言和策略参考，不能成为患者画像证据。",
    )
    add_heading(doc, "知识范围", 2)
    add_simple_table(
        doc,
        ["知识类别", "必须覆盖的内容", "主要检索条件"],
        [
            ("MI阶段", "参与 聚焦 唤出 计划四个阶段的目标和语言范式", "当前阶段和阶段目标"),
            ("OARS技术", "开放式提问 肯定 简单反映 复杂反映 总结", "允许技术和患者表达"),
            ("症状探索", "抑郁 焦虑 创伤 强迫 失眠 躁狂史 精神病性症状 物质和躯体影响", "症状领域和缺失画像字段"),
            ("量表引入", "PHQ 9 GAD 7 PCL 5 Y BOCS ISI MDQ的解释 同意 拒绝和中断", "量表领域和许可状态"),
            ("目标与动机", "患者目标 改变重要性 信心 矛盾心理和阻力", "MI阶段和动机水平"),
            ("错误范式", "机械复述 连续多问 过早量表化 提前诊断和给建议", "输出校验失败类型"),
        ],
        [Inches(1.2), Inches(3.7), Inches(1.85)],
    )

    add_heading(doc, "知识卡字段", 2)
    add_body(
        doc,
        "每条知识卡至少记录knowledge id 来源 许可 语言 MI阶段 OARS技术 症状领域 对话目的 患者表达模式 待补字段 推荐示例 反例 使用理由 量表领域 禁忌 标签 版本和人工审核状态。来源和许可字段必须在导入索引前完成。",
    )
    add_heading(doc, "接入现有L2的流程", 2)
    for text in (
        "L2语义模型从患者消息中提取症状 情绪 意图和可能量表领域。",
        "规则读取当前MI阶段 待补字段 L1安全命令和量表许可状态。",
        "L2根据结构化状态构造检索请求，而不是只用患者原句检索。",
        "通用RAG返回少量互补范式，例如阶段范式 症状范式和OARS范式各一条。",
        "模型结合范式生成自然回复，不能复制示例中的患者事实。",
        "现有MI回复校验 量表规则 信息充分度和出口守卫继续执行。",
        "检索knowledge id写入调试记录，但不能写入患者evidence nodes。",
    ):
        add_number(doc, text)

    add_heading(doc, "任务清单", 2)
    add_task_table(doc, [
        ("L2K1", "来源清单", "寻找公开MI 问诊 量表沟通和心理访谈资料并记录许可", "来源与许可清单", "每条资料可追溯且允许当前用途"),
        ("L2K2", "标签体系", "定义MI阶段 OARS技术 症状领域 对话意图和禁忌标签", "标签词典和映射表", "同一标签含义唯一 覆盖现有L2字段"),
        ("L2K3", "知识卡加工", "把原始资料整理成推荐范式 反例和适用条件", "标准化知识卡文件", "无患者隐私 无诊断越界 有人工审核状态"),
        ("L2K4", "术语库", "整理症状 情绪 认知模式 患者目标和改变动机表达", "术语及同义表达库", "只帮助理解 不直接写入患者事实"),
        ("L2K5", "量表范式", "整理六个量表的解释 同意 拒绝 中断和自然穿插范式", "量表沟通知识卡", "不改写题干 不绕过患者许可"),
        ("L2K6", "检索构造", "使用阶段 待补字段 症状领域和许可状态生成查询", "L2检索适配器", "查询不包含未经确认的推断字段"),
        ("L2K7", "Prompt接入", "将Top K范式作为参考加入现有L2调用", "Prompt组装调整", "检索失败时可按现有流程继续且不伪造知识"),
        ("L2K8", "质量验证", "验证MI风格 单问题优先 量表时机和禁止诊断", "测试集和人工抽检报告", "回复自然且不因RAG破坏规则控制"),
    ])

    add_heading(doc, "验收场景", 2)
    for text in (
        "患者首次表达压力时，系统优先建立关系，不立即触发量表。",
        "患者提及失眠后，系统能检索睡眠探索范式并提出一个自然问题。",
        "建议PHQ 9前能说明用途 时长和非诊断性质，并征得同意。",
        "患者拒绝量表时继续临床访谈，不重复强迫邀请。",
        "知识库示例中的人物 症状和分数不会进入当前患者画像。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "负责人交付边界", 2)
    add_bullet(doc, "负责人  MI资料 数据和L2开发成员；接入点为layer2 service dialogue llm client和mi response validator。")
    add_bullet(doc, "不得修改  L1安全决策 量表题目和分值规则 L2到L3最低字段要求。")

    add_heading(doc, "三  L3临床诊断知识库", 1)
    add_body(
        doc,
        "L3知识库为CLINICR提供外部诊断标准、排除条件和鉴别诊断依据。它不替代当前候选生成、逐条排除和Verifier，只为这些步骤补充有版本和来源的临床知识。患者证据与临床知识必须分别保存，不能相互替代。",
    )
    add_heading(doc, "知识范围", 2)
    add_simple_table(
        doc,
        ["知识类别", "必须记录的内容", "CLINICR用途"],
        [
            ("疾病概览", "名称 编码 核心表现 病程 功能损害和主要排除条件", "Step 1同时考虑规定的10种障碍"),
            ("诊断标准", "条目代码 名称 标准内容 核心性 数量和组合要求", "Step 2逐条核对"),
            ("排除条件", "物质 药物 躯体疾病 躁狂史 精神病性症状等", "阻止不满足核心条件的候选通过"),
            ("鉴别诊断", "共同症状 关键差异和补充问题", "比较候选并返回L2补充信息"),
            ("严重度和转介", "严重度依据 安全条件和专业转介提示", "疗法匹配前的规则输入"),
            ("来源信息", "机构 文件 版本 章节 页码 网址 许可和生效日期", "建立知识证据链和更新管理"),
        ],
        [Inches(1.25), Inches(3.65), Inches(1.85)],
    )

    add_heading(doc, "双证据链要求", 2)
    add_body(
        doc,
        "每条标准判断必须同时包含patient evidence refs和knowledge refs。前者指向L2记录的患者原话或量表答案，后者指向临床知识库中的标准条目。只有患者证据不能证明采用了正确标准，只有知识条目也不能证明患者满足该标准。",
    )
    add_simple_table(
        doc,
        ["证据类型", "回答的问题", "示例", "来源"],
        [
            ("患者证据", "患者实际报告了什么", "EV 023  最近一个多月整个人都没劲", "L2对话和证据节点"),
            ("临床知识", "这项表现对应哪条标准", "CLINICAL MDD A1 001", "官方或合法授权资料"),
            ("条目结论", "为什么标为符合 不符合或待确认", "A1 met并引用两类证据", "L3结构化判断"),
        ],
        [Inches(1.15), Inches(2.05), Inches(2.2), Inches(1.35)],
    )

    add_heading(doc, "接入现有CLINICR的流程", 2)
    for text in (
        "Step 1使用10种障碍的受控概览知识，模型仍需完整考虑规定范围。",
        "Step 1生成3至5个候选后，按候选疾病代码和资料版本精确检索完整标准包。",
        "每个候选标准包包含诊断条目 病程 功能损害 排除和鉴别信息。",
        "Step 2只允许引用本次标准包中的knowledge refs和患者画像白名单中的evidence refs。",
        "规则验证条目覆盖 数量门槛 核心条件 引用存在性和版本一致性。",
        "Verifier独立检查证据充分性 逻辑连贯性 鉴别覆盖和推理跳跃。",
        "缺少患者证据时返回L2，临床知识库不可用时报告知识检索错误，不能依赖模型记忆补齐。",
    ):
        add_number(doc, text)

    add_heading(doc, "任务清单", 2)
    add_task_table(doc, [
        ("L3K1", "来源与授权", "确定官方分类 临床指南 专业机构资料和允许使用的范围", "来源 版本和许可登记", "无来源或无授权资料不进入知识库"),
        ("L3K2", "疾病映射", "为当前10种障碍建立标准名称 编码 同义词和内部标签", "疾病映射表", "名称和编码唯一 可用于精确过滤"),
        ("L3K3", "疾病概览", "整理核心表现 病程 功能要求和主要排除项", "10种障碍概览卡", "Step 1能完整加载且版本一致"),
        ("L3K4", "标准条目", "按疾病和标准代码拆分条目并标记核心性 数量和组合要求", "结构化ClinicalCriterion", "条目代码完整且有来源位置"),
        ("L3K5", "鉴别知识", "整理候选间共同表现 关键差异和建议补充字段", "鉴别诊断知识卡", "能映射到L2可收集字段"),
        ("L3K6", "标准包构建", "按候选疾病生成固定版本的Clinical Knowledge Pack", "标准包服务", "一次诊断过程不混用不同版本"),
        ("L3K7", "CLINICR接入", "把概览接入Step 1 完整标准接入Step 2", "L3知识适配器", "保留候选数量 门槛 回环和Verifier独立性"),
        ("L3K8", "双引用守卫", "验证patient evidence refs和knowledge refs均来自白名单", "L3输出校验规则", "未知引用 缺失来源和虚构标准均被拒绝"),
        ("L3K9", "临床抽检", "按疾病检查标准完整性和结论可解释性", "覆盖矩阵和审核记录", "10种障碍均通过资料和流程验收"),
    ])

    add_heading(doc, "负责人交付边界", 2)
    add_bullet(doc, "建议负责人  L3开发成员和临床资料审核成员共同负责。")
    add_bullet(doc, "主要接入点  layers layer3 diagnosis clinicr llm client api和output validator。")
    add_bullet(doc, "不得修改  CLINICR三步顺序 候选置信度门槛 Verifier独立评分和L3到L2回环原则。")
    add_bullet(doc, "资料要求  DSM相关材料必须确认版权和授权，优先使用公开可用的ICD及权威临床指南。")

    add_heading(doc, "四  Patient Agent", 1)
    add_body(
        doc,
        "Patient Agent用于开发和测试环境，模拟具有固定背景、症状和行为模式的患者，与完整系统进行多轮交互。它帮助团队复现复杂路径和自动运行疗程，但不能代替真实患者评估，也不能进入生产患者流程。",
    )
    add_heading(doc, "病例模型", 2)
    add_simple_table(
        doc,
        ["字段组", "内容", "用途"],
        [
            ("基本背景", "年龄段 工作 家庭关系和主要事件", "保持病例上下文一致"),
            ("临床事实", "症状 病程 功能影响 量表潜在答案和排除条件", "回答L2并验证L3"),
            ("心理特征", "情绪 认知模式 核心信念 目标和改变动机", "模拟MI和治疗响应"),
            ("风险脚本", "无风险 模糊风险 明确意念 计划和即时危机", "覆盖L1安全路径"),
            ("披露策略", "主动信息 隐藏信息 需要追问的信息和拒答条件", "避免首次回答泄露全部病例"),
            ("行为参数", "合作程度 表达长度 情绪强度 反复和矛盾", "生成不同类型患者"),
            ("变化轨迹", "改善 无变化 恶化 复发和退出", "测试疗程推进与终止"),
        ],
        [Inches(1.2), Inches(3.75), Inches(1.8)],
    )

    add_heading(doc, "运行流程", 2)
    for text in (
        "测试控制器加载固定PatientScenario和随机种子。",
        "Patient Agent根据病例隐藏状态生成首条患者消息。",
        "完整系统从L1开始处理并返回患者可见回复。",
        "Patient Agent只根据病例事实 对话历史和披露策略回答。",
        "测试控制器记录每轮输入 路由 输出 状态和预期事件。",
        "达到终止条件 危机转介 人工复核或最大测试轮数后结束。",
        "评测器检查事实一致性 路径覆盖 证据完整性和终止结果。",
    ):
        add_number(doc, text)

    add_heading(doc, "任务清单", 2)
    add_task_table(doc, [
        ("PA1", "PatientScenario", "定义背景 临床事实 风险 披露策略和变化轨迹", "病例Schema和示例", "病例字段完整且不依赖自由文本猜测"),
        ("PA2", "病例库", "建设低风险 高风险 信息不足 共病和诊断排除等病例", "版本化测试病例", "每个病例有预期路径和关键事实"),
        ("PA3", "患者生成器", "根据隐藏状态和历史生成自然患者回复", "Patient Agent模型适配器", "不泄露系统提示 不添加病例外事实"),
        ("PA4", "一致性守卫", "校验病程 量表 风险和背景前后一致", "事实一致性检查器", "矛盾输出被记录并终止该次评测"),
        ("PA5", "自动对话控制", "循环调用完整系统API并处理确认 暂停 量表和治疗", "场景运行器", "能自动完成评估 诊断 治疗和结案路径"),
        ("PA6", "停止条件", "配置最大轮数 目标事件 危机 人工复核和系统终止", "测试级终止控制", "测试不会无限运行且不冒充疗程终止器"),
        ("PA7", "评测报告", "统计路径覆盖 错误类型 事实一致性和模型耗时", "结构化报告", "失败可定位到层级和具体轮次"),
    ])

    add_heading(doc, "验收场景", 2)
    for text in (
        "低风险抑郁样例完成L2画像确认并进入L3。",
        "信息不足样例触发L3返回L2并能补充指定字段。",
        "明确危机样例在L1进入安全路径且不调用后续层。",
        "候选诊断被排除时，Patient Agent保持原病例事实不迎合系统。",
        "治疗改善 无改善和主动退出三种轨迹能够触发不同终止结果。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "负责人交付边界", 2)
    add_bullet(doc, "负责人  智能体和自动化测试成员；建议目录为simulation patient agent和tests scenarios。")
    add_bullet(doc, "不得修改  生产环境患者处理逻辑 诊断标准和治疗决策门槛。")

    add_heading(doc, "五  前端和完整系统API", 1)
    add_body(
        doc,
        "前端应通过一个完整系统API调用顶层编排器，不分别调用L1至L5。现有main.py是终端入口，现有L4 API只覆盖治疗层，均不能单独承担完整前端后端接口。前端工作因此包含完整系统API封装和两类界面。",
    )
    add_heading(doc, "角色和界面", 2)
    add_simple_table(
        doc,
        ["界面", "主要功能", "禁止展示"],
        [
            ("患者端", "自然语言对话 量表交互 画像确认 暂停恢复 安全提示和疗程结束信息", "候选诊断 内部评分 原始Prompt和开发调试信息"),
            ("开发者端", "当前路由 L1风险 L2画像 L3证据 L4指令 L5审核 RAG命中和Agent状态", "API密钥 未脱敏个人信息和无权限会话"),
            ("演示控制台", "选择测试病例 启动Patient Agent 查看流程和导出结果", "生产患者操作入口"),
        ],
        [Inches(1.25), Inches(3.85), Inches(1.65)],
    )

    add_heading(doc, "完整系统API", 2)
    add_simple_table(
        doc,
        ["接口", "用途", "主要返回"],
        [
            ("创建会话", "建立顶层会话并返回初始状态", "session id 路由和初始提示"),
            ("发送消息", "把患者消息交给顶层编排器", "患者回复 等待类型 路由和错误码"),
            ("画像确认", "提交继续 修改或暂停决定", "下一步状态和患者可见回复"),
            ("提交量表", "提交当前许可题目的答案", "下一题或返回访谈"),
            ("暂停恢复", "控制患者主动暂停状态", "当前可执行动作"),
            ("获取开发状态", "读取开发者可见的流程快照", "各层结构化状态和证据"),
            ("获取疗程总结", "读取L5生成的周期总结", "EpisodeSummary"),
            ("获取历史疗程", "读取外部Memory", "脱敏疗程列表和版本"),
        ],
        [Inches(1.35), Inches(3.45), Inches(1.95)],
    )

    add_heading(doc, "任务清单", 2)
    add_task_table(doc, [
        ("FE1", "用户流程", "绘制患者评估 画像确认 诊断等待 治疗 安全和结案流程", "页面流程图和状态表", "每个编排状态有对应页面和可执行操作"),
        ("FE2", "系统API", "用Web服务封装FiveLayerOrchestrator和会话操作", "完整系统后端接口", "前端不直接调用单层且错误码稳定"),
        ("FE3", "患者聊天", "实现消息 处理中状态 量表卡 确认卡 暂停和结束", "患者端Demo", "患者看不到开发者字段且异常有明确提示"),
        ("FE4", "开发者面板", "展示路由 画像 候选诊断 双证据链 指令 审核和RAG结果", "可视化调试面板", "数据来自结构化状态而非解析终端文本"),
        ("FE5", "Agent控制台", "选择病例 启动停止自动对话和查看逐轮记录", "测试演示页面", "与真实患者入口明确隔离"),
        ("FE6", "错误状态", "区分API错误 超时 Schema错误 知识检索错误 人工复核和会话停止", "错误映射表和页面", "错误不会显示密钥或内部堆栈"),
        ("FE7", "会话与Memory", "刷新后按会话ID恢复可持久化状态并查看历史", "恢复和历史界面", "不能从开发者画像HTML恢复会话"),
        ("FE8", "联调与演示", "完成手工患者和Patient Agent两条演示路径", "可运行Demo和演示脚本", "从创建会话到疗程总结可完整操作"),
    ])

    add_heading(doc, "负责人交付边界", 2)
    add_bullet(doc, "建议负责人  前端成员负责页面，后端接口成员负责完整系统API。")
    add_bullet(doc, "主要接入点  orchestration orchestrator session store和新增系统级API目录。")
    add_bullet(doc, "不得修改  前端不得自行计算风险 诊断 量表分数 疗法或审核结果。")

    doc.add_page_break()
    add_heading(doc, "六  联调顺序和里程碑", 1)
    add_simple_table(
        doc,
        ["里程碑", "并行工作", "完成条件"],
        [
            ("M1 接口冻结", "五项负责人共同确认Schema和文件边界", "接口文档完成且Schema一致性测试通过"),
            ("M2 数据可用", "L2和L3分别完成首版知识数据", "来源 许可 标签和人工审核字段完整"),
            ("M3 检索可用", "RAG完成两库隔离检索", "固定查询集能返回预期Top K和来源"),
            ("M4 层级接入", "L2和L3分别接入RAG", "原有规则和守卫仍通过 检索失败可定位"),
            ("M5 自动模拟", "Patient Agent调用完整系统API", "能完成评估 回环 治疗和终止场景"),
            ("M6 前端联调", "患者端 开发者端和Agent控制台接入", "完整流程不依赖终端文本"),
            ("M7 项目验收", "真实LLM 受控模型和人工抽检", "功能 流程 证据 错误和演示清单全部通过"),
        ],
        [Inches(1.25), Inches(3.45), Inches(2.05)],
    )

    add_heading(doc, "推荐并行安排", 2)
    add_simple_table(
        doc,
        ["角色", "主任务", "需要配合的成员"],
        [
            ("项目和接口负责人", "统一Schema 文件边界 合并和里程碑验收", "全部成员"),
            ("RAG负责人", "R1至R8", "L2知识 L3知识和后端成员"),
            ("L2知识负责人", "L2K1至L2K8", "数据审核 RAG和L2成员"),
            ("L3知识负责人", "L3K1至L3K9", "临床审核 RAG和L3成员"),
            ("Patient Agent负责人", "PA1至PA7", "测试 后端和前端成员"),
            ("前端负责人", "FE1 FE3至FE8", "系统API和各层负责人"),
            ("系统API负责人", "FE2及鉴权 错误和会话接口", "编排器 前端和Memory成员"),
            ("测试负责人", "固定查询集 场景集 集成测试和验收记录", "全部成员"),
        ],
        [Inches(1.55), Inches(3.2), Inches(2.0)],
    )

    add_heading(doc, "七  全局验收清单", 1)
    checks = [
        "L2和L3使用不同collection，不发生知识混检。",
        "所有检索结果包含knowledge id 来源 版本和相关性分数。",
        "L2知识示例不会进入患者证据节点，也不会提前生成诊断。",
        "L3每条关键诊断判断同时引用患者证据和临床知识。",
        "RAG不可用时能区分检索错误 API错误和Schema错误。",
        "Patient Agent不会添加病例外事实，测试能够按条件停止。",
        "患者端只展示患者可见内容，开发者端不泄露API密钥。",
        "前端只调用完整系统API，不逐层拼装L1至L5。",
        "原有L1至L5自动化测试继续通过。",
        "新增或修改Pydantic模型后，Schema参考文档已经同步生成并通过一致性测试。",
        "真实LLM测试和受控模型测试分别记录，不能用受控模型结果代替真实API结论。",
        "项目演示能够从创建会话运行到疗程终止和总结查看。",
    ]
    for text in checks:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.first_line_indent = Inches(-0.1)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("□  " + text)
        set_run_fonts(r)

    add_heading(doc, "八  成员提交规范", 1)
    add_body(doc, "每位负责人提交任务时，应同时提交代码或数据、接口说明、测试结果和未解决问题。提交记录至少包含以下内容。")
    for text in (
        "任务编号和负责人。",
        "本次新增或修改的文件。",
        "输入 输出和错误类型。",
        "自动化测试命令和结果。",
        "真实LLM测试是否执行及使用的模型配置。",
        "已知限制和需要其他成员配合的接口。",
        "下一阶段计划和预计完成时间。",
    ):
        add_bullet(doc, text)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
